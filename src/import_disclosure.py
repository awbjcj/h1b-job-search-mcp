"""One-shot cache conversion; the serving process never imports pandas/openpyxl."""

import argparse
import math
import os
import sqlite3
import tempfile
from datetime import date, datetime
from pathlib import Path

from disclosure_store import SCHEMA_VERSION, WAGE_COLUMNS, configure, numeric, quote
from employers import _normalise_employer
from file_cache import release_file_cache

BATCH_SIZE = 1000


def scalar(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def write_database(destination, columns, rows):
    """Build beside the target and atomically publish only complete databases."""
    columns = [str(c) for c in columns]
    if len(set(c.casefold() for c in columns)) != len(columns):
        raise ValueError("Duplicate disclosure columns")
    if any(c.startswith("_h1b_") for c in columns):
        raise ValueError("Reserved disclosure column name")
    employer = next(
        (c for c in ("EMPLOYER_NAME", "EMPLOYER_BUSINESS_DBA") if c in columns), None
    )
    if employer is None:
        raise ValueError("Disclosure has no employer column")
    employer_index = columns.index(employer)
    wage = next((c for c in WAGE_COLUMNS if c in columns), None)
    wage_index = columns.index(wage) if wage else None
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(
        prefix=destination.name + ".", suffix=".building", dir=destination.parent
    )
    os.close(fd)
    try:
        db = sqlite3.connect(temporary)
        try:
            configure(db)
            db.execute(
                "CREATE TABLE disclosures ("
                + ", ".join(map(quote, columns))
                + ", _h1b_employer_key TEXT, _h1b_wage REAL)"
            )
            sql = (
                "INSERT INTO disclosures VALUES ("
                + ",".join("?" for _ in range(len(columns) + 2))
                + ")"
            )
            batch, count = [], 0
            for row in rows:
                values = tuple(scalar(value) for value in row)
                if len(values) != len(columns):
                    raise ValueError("Disclosure row width does not match header")
                key = _normalise_employer(values[employer_index])
                wage_value = (
                    numeric(values[wage_index]) if wage_index is not None else None
                )
                batch.append((*values, key, wage_value))
                if len(batch) >= BATCH_SIZE:
                    db.executemany(sql, batch)
                    count += len(batch)
                    batch.clear()
            db.executemany(sql, batch)
            count += len(batch)
            if not count:
                raise ValueError("Disclosure contains no rows")
            db.execute("CREATE INDEX employer_key ON disclosures (_h1b_employer_key)")
            db.execute(f"CREATE INDEX employer_name ON disclosures ({quote(employer)})")
            db.execute("CREATE TABLE metadata (row_count INTEGER NOT NULL)")
            db.execute("INSERT INTO metadata VALUES (?)", (count,))
            db.execute(f"PRAGMA user_version={SCHEMA_VERSION}")
            db.commit()
            if db.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                raise ValueError("Disclosure cache failed integrity check")
        finally:
            db.close()
            release_file_cache(temporary)
        os.replace(temporary, destination)
        return count
    finally:
        for suffix in ("", "-journal"):
            Path(temporary + suffix).unlink(missing_ok=True)


def convert(source, destination):
    try:
        return _convert(source, destination)
    finally:
        release_file_cache(source)


def _convert(source, destination):
    source = Path(source)
    if source.suffix == ".pkl":
        # Only locally generated, trusted legacy caches are accepted here.
        # Pickle cannot stream; its one-time allocation dies with this process.
        import pandas as pd

        frame = pd.read_pickle(source)
        if not isinstance(frame, pd.DataFrame):
            raise TypeError("Cached H-1B data is not a DataFrame")
        rows = (
            tuple(None if value is pd.NA or value is pd.NaT else value for value in row)
            for row in frame.itertuples(index=False, name=None)
        )
        return write_database(destination, frame.columns, rows)
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        # Do not silently truncate files whose declared dimensions are wrong.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        columns = [
            str(value) if value is not None else f"Unnamed: {i}"
            for i, value in enumerate(header)
        ]

        def padded_rows():
            for row in rows:
                if all(value is None for value in row):
                    continue
                yield (*row, *([None] * (len(columns) - len(row))))

        return write_database(destination, columns, padded_rows())
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("destination")
    args = parser.parse_args()
    print(f"Indexed {convert(args.source, args.destination)} disclosure rows")
